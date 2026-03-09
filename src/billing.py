from __future__ import annotations

import datetime as dt
import json
import re
import uuid
from pathlib import Path
from typing import List, Optional, Tuple

import pandas as pd
import openpyxl


# -----------------------------
# Storage
# -----------------------------

def _data_dir() -> Path:
    base = Path(__file__).resolve().parent.parent
    d = base / "data" / "facturation"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _records_path() -> Path:
    return _data_dir() / "records.csv"


def _meta_path() -> Path:
    return _data_dir() / "meta.json"


def _read_meta() -> dict:
    p = _meta_path()
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _write_meta(meta: dict) -> None:
    _meta_path().write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


# -----------------------------
# Helpers
# -----------------------------

def _norm(s: str) -> str:
    return (str(s or "")).strip().lower()


def norm_site_facturation(site: str) -> str:
    s0 = str(site or "").strip()
    sN = _norm(s0)

    if re.fullmatch(r"24\s*(ter|simple)", sN):
        return "Internat"
    if sN in {"24ter", "24simple"}:
        return "Internat"
    if "24" in sN and ("ter" in sN or "simple" in sN):
        return "Internat"

    return s0


def is_pdj_regime(regime: str) -> bool:
    r = _norm(regime)
    return ("pdj" in r) or ("petit" in r and "dej" in r)


def is_mixe_lisse_regime(regime: str) -> bool:
    r = _norm(regime)
    return ("mixe" in r) or ("mixé" in r) or ("lisse" in r)


# -----------------------------
# DATE FIX (BUG CORRECTION)
# -----------------------------

def _date_from_week_and_dayname(week_monday: dt.date, day_name: str) -> dt.date:
    day_offsets = {
        "Lundi": 0,
        "Mardi": 1,
        "Mercredi": 2,
        "Jeudi": 3,
        "Vendredi": 4,
        "Samedi": 5,
        "Dimanche": 6,
    }

    name = str(day_name).strip()

    if name not in day_offsets:
        raise ValueError(f"Jour invalide dans planning: {day_name}")

    return week_monday + dt.timedelta(days=day_offsets[name])


# -----------------------------
# Planning → daily totals
# -----------------------------

def planning_to_daily_totals(
    dej: pd.DataFrame,
    din: pd.DataFrame,
    week_monday: dt.date,
) -> pd.DataFrame:

    df = pd.concat([dej, din], ignore_index=True)

    day_cols = [
        c for c in
        ["Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"]
        if c in df.columns
    ]

    melted = df.melt(
        id_vars=["Site"],
        value_vars=day_cols,
        var_name="day_name",
        value_name="qty",
    )

    melted["qty"] = pd.to_numeric(melted["qty"], errors="coerce").fillna(0).astype(int)

    melted["date"] = melted["day_name"].map(
        lambda d: _date_from_week_and_dayname(week_monday, d)
    )

    out = melted.groupby(["date", "Site"], as_index=False)["qty"].sum()

    out = out.rename(columns={"Site": "site", "qty": "qty_repas"})

    return out


def mixe_lisse_to_daily_totals(
    dej: Optional[pd.DataFrame],
    din: Optional[pd.DataFrame],
    week_monday: dt.date,
) -> pd.DataFrame:

    frames = []
    for df in (dej, din):
        if df is not None and not df.empty:
            frames.append(df.copy())

    if not frames:
        return pd.DataFrame(columns=["date","site","qty_ml"])

    df = pd.concat(frames, ignore_index=True)

    day_cols = [
        c for c in
        ["Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"]
        if c in df.columns
    ]

    melted = df.melt(
        id_vars=["Site"],
        value_vars=day_cols,
        var_name="day_name",
        value_name="qty",
    )

    melted["qty"] = pd.to_numeric(melted["qty"], errors="coerce").fillna(0).astype(int)

    melted["date"] = melted["day_name"].map(
        lambda d: _date_from_week_and_dayname(week_monday, d)
    )

    out = melted.groupby(["date","Site"], as_index=False)["qty"].sum()

    out = out.rename(columns={"Site":"site","qty":"qty_ml"})

    return out


# -----------------------------
# Persist / load
# -----------------------------

def save_week(
    *,
    week_monday: dt.date,
    repas_daily: pd.DataFrame,
    ml_daily: pd.DataFrame,
    source_filename: str = "",
) -> Tuple[int, int]:

    def _norm_df(df: pd.DataFrame, qty_col: str, category: str):

        if df is None or df.empty:
            return pd.DataFrame()

        out = df.copy()

        out["date"] = pd.to_datetime(out["date"]).dt.date
        out["site"] = out["site"].astype(str)
        out["category"] = category
        out["qty"] = out[qty_col]
        out["week_monday"] = week_monday.isoformat()
        out["source"] = source_filename
        out["record_id"] = [str(uuid.uuid4()) for _ in range(len(out))]

        return out[
            ["record_id","date","site","category","qty","week_monday","source"]
        ]

    a = _norm_df(repas_daily,"qty_repas","repas")
    b = _norm_df(ml_daily,"qty_ml","mixe_lisse")

    new = pd.concat([a,b], ignore_index=True)

    p = _records_path()

    if p.exists():
        old = pd.read_csv(p, parse_dates=["date"])
    else:
        old = pd.DataFrame()

    merged = pd.concat([old,new], ignore_index=True)

    merged.to_csv(p, index=False)

    return len(a), len(b)


def load_records() -> pd.DataFrame:

    p = _records_path()

    if not p.exists():
        return pd.DataFrame()

    df = pd.read_csv(p, parse_dates=["date"])

    return df


# -----------------------------
# Delete
# -----------------------------

def delete_billing_records(
    *,
    week_monday: Optional[dt.date] = None,
    record_ids: Optional[List[str]] = None,
) -> int:

    p = _records_path()

    if not p.exists():
        return 0

    df = pd.read_csv(p)

    before = len(df)

    if week_monday:
        df = df[df["week_monday"] != week_monday.isoformat()]

    if record_ids:
        df = df[~df["record_id"].isin(record_ids)]

    df.to_csv(p, index=False)

    return before - len(df)


# -----------------------------
# Export
# -----------------------------

def export_monthly_workbook(
    records: pd.DataFrame,
    out_path: str,
) -> str:

    if records.empty:
        wb = openpyxl.Workbook()
        wb.save(out_path)
        return out_path

    records["date"] = pd.to_datetime(records["date"])

    records["year"] = records["date"].dt.year
    records["month"] = records["date"].dt.month

    year = records["year"].max()

    wb = openpyxl.Workbook()

    for m in range(1,13):

        ws = wb.create_sheet(f"{year}-{m:02d}")

        sub = records[(records["year"]==year)&(records["month"]==m)]

        if sub.empty:
            continue

        pivot = sub.pivot_table(
            index="date",
            columns="site",
            values="qty",
            aggfunc="sum",
            fill_value=0,
        )

        for r,row in enumerate(pivot.itertuples(),start=1):

            ws.cell(r,1,str(row.Index.date()))

            for c,val in enumerate(row[1:],start=2):

                ws.cell(r,c,val)

    wb.save(out_path)

    return out_path


# -----------------------------
# Import corrected workbook
# -----------------------------

def apply_corrected_monthly_workbook(
    xlsx_path: str,
) -> Tuple[int,int]:

    df = pd.read_excel(xlsx_path)

    p = _records_path()

    if p.exists():
        old = pd.read_csv(p)
    else:
        old = pd.DataFrame()

    old = df

    old.to_csv(p,index=False)

    return 0,len(df)